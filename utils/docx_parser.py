import re
from typing import List, Dict, Optional
from dataclasses import dataclass, field


@dataclass
class FieldInfo:
    name: str
    type: str
    desc: str
    nullable: bool = True


@dataclass
class Relation:
    target_table: str
    relation_type: str
    via_field: str
    via: str


@dataclass
class TableInfo:
    table_name: str
    table_type: str
    table_desc: str
    business_rules: List[str] = field(default_factory=list)
    fields: List[FieldInfo] = field(default_factory=list)
    relations: List[Relation] = field(default_factory=list)


class DocxParser:
    def __init__(self, file_path: str):
        self.file_path = file_path
        self.paragraphs: List[str] = []
        self.tables: List[TableInfo] = []

    def parse(self) -> List[TableInfo]:
        from docx import Document

        doc = Document(self.file_path)
        self.paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        i = 0
        while i < len(self.paragraphs):
            para = self.paragraphs[i]

            if self._is_table_header(para):
                table_info = self._parse_table(i)
                if table_info:
                    self.tables.append(table_info)
                    i = self._skip_to_next_table(i)
                    continue

            i += 1

        return self.tables

    def _is_table_header(self, para: str) -> bool:
        pattern = r"^\d+\.\s*表名：\S+（\S+表）"
        return bool(re.match(pattern, para))

    def _parse_table(self, start_idx: int) -> Optional[TableInfo]:
        para = self.paragraphs[start_idx]

        table_match = re.match(r"^(\d+)\.\s*表名：(\S+)（(\S+)）", para)
        if not table_match:
            return None

        table_name = table_match.group(2)
        table_type = table_match.group(3)

        table_info = TableInfo(
            table_name=table_name,
            table_type=table_type,
            table_desc=""
        )

        i = start_idx + 1
        section = None

        while i < len(self.paragraphs):
            para = self.paragraphs[i]

            if self._is_table_header(para):
                break

            if para.startswith("表描述："):
                desc = para.replace("表描述：", "").strip()
                if "业务规则" in desc:
                    parts = desc.split("业务规则")
                    table_info.table_desc = parts[0].strip()
                else:
                    table_info.table_desc = desc

            elif para.startswith("业务规则："):
                section = "business_rules"

            elif "字段详情" in para or para.startswith("字段详情"):
                section = "fields"

            elif para.startswith("表关系："):
                section = "relations"

            elif section == "business_rules" and para and not para.startswith("字段详情") and not para.startswith("表关系"):
                rule = para.lstrip("- ").strip()
                if rule:
                    table_info.business_rules.append(rule)

            elif section == "fields" and not para.startswith("表关系") and not para.startswith("业务规则"):
                field_info = self._parse_field_line(para)
                if field_info:
                    table_info.fields.append(field_info)

            elif section == "relations" and para.startswith("与"):
                relation = self._parse_relation_line(para)
                if relation:
                    table_info.relations.append(relation)

            i += 1

        return table_info

    def _parse_field_line(self, line: str) -> Optional[FieldInfo]:
        line = line.strip()

        match = re.match(r"([^：]+)：(.+)", line)
        if not match:
            return None

        field_name = match.group(1).strip()
        rest = match.group(2).strip()

        type_match = re.match(r"([^（]+)（([^）]+)）", rest)
        if type_match:
            field_type = type_match.group(1).strip()
            desc = type_match.group(2).strip()
        else:
            field_type = "unknown"
            desc = rest

        nullable = "不可为空" not in desc

        return FieldInfo(
            name=field_name,
            type=field_type,
            desc=desc,
            nullable=nullable
        )

    def _parse_relation_line(self, line: str) -> Optional[Relation]:
        via_match = re.search(r"（([^）]+)）", line)
        via = via_match.group(1) if via_match else ""

        table_match = re.search(r"与\s*(\S+)\s*表", line)
        target_table = table_match.group(1) if table_match else ""

        if "多对一" in line:
            relation_type = "多对一"
        elif "一对多" in line:
            relation_type = "一对多"
        elif "多对多" in line:
            relation_type = "多对多"
        elif "一对一" in line:
            relation_type = "一对一"
        else:
            relation_type = "未知"

        via_field = ""
        if "→" in via:
            parts = via.split("→")
            if len(parts) == 2:
                from_part = parts[0].strip()
                if "." in from_part:
                    via_field = from_part.split(".")[-1]
                elif from_part:
                    via_field = from_part

        return Relation(
            target_table=target_table,
            relation_type=relation_type,
            via_field=via_field,
            via=via
        )

    def _skip_to_next_table(self, current_idx: int) -> int:
        for i in range(current_idx + 1, len(self.paragraphs)):
            if self._is_table_header(self.paragraphs[i]):
                return i
        return len(self.paragraphs)


def parse_docx(file_path: str) -> List[TableInfo]:
    parser = DocxParser(file_path)
    return parser.parse()


if __name__ == "__main__":
    tables = parse_docx("resource/htower.docx")
    for table in tables:
        print(f"\n{'='*60}")
        print(f"表名: {table.table_name} ({table.table_type})")
        print(f"字段数: {len(table.fields)}, 关系数: {len(table.relations)}")
        for rel in table.relations:
            print(f"  → {rel.target_table} ({rel.relation_type}) via {rel.via_field}")
